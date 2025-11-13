import os
import re
import json
import pandas as pd
from collections import defaultdict
from email import policy
from email.parser import BytesParser
from pathlib import Path

try:
    import docx
except Exception:
    docx = None
try:
    import PyPDF2
except Exception:
    PyPDF2 = None

OUTPUT_DIR = "output"
REPORT_PATH = os.path.join(OUTPUT_DIR, "validation_report.txt")

def load_config():
    with open("config.json", "r", encoding="utf-8") as f:
        return json.load(f)

def load_mapping():
    xlsx = os.path.join("output", "mapping_final.xlsx")
    csvf = os.path.join("output", "mapping_meta.csv")
    if os.path.exists(xlsx):
        df = pd.read_excel(xlsx, dtype=str)
    elif os.path.exists(csvf):
        df = pd.read_csv(csvf, dtype=str)
    else:
        raise FileNotFoundError("mapping_final.xlsx or mapping_meta.csv not found in output/")
    if 'doc_id' in df.columns:
        try:
            df['doc_id'] = df['doc_id'].astype(int)
        except Exception:
            pass
    return df

def build_sit_regex_map(config):
    sit_map = {}
    for s in config.get("sits", []):
        sid = s["id"]
        regex = s.get("regex", "")
        if regex:
            try:
                sit_map[sid] = re.compile(regex, flags=re.MULTILINE | re.IGNORECASE)
            except re.error:
                sit_map[sid] = None
        else:
            sit_map[sid] = None
    return sit_map

def read_txt(path):
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""

def read_docx(path):
    if docx is None:
        return ""
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        return "\n".join(parts)
    except Exception:
        return ""

def read_pdf(path):
    if PyPDF2 is None:
        return ""
    try:
        out = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for p in reader.pages:
                try:
                    out.append(p.extract_text() or "")
                except Exception:
                    continue
        return "\n".join(out)
    except Exception:
        return ""

def read_eml(path):
    try:
        with open(path, "rb") as f:
            msg = BytesParser(policy=policy.default).parse(f)
        if msg.is_multipart():
            parts = []
            for part in msg.walk():
                ctype = part.get_content_type()
                if ctype == "text/plain":
                    parts.append(part.get_content())
            return "\n".join(parts)
        else:
            return msg.get_content()
    except Exception:
        return ""

def extract_text_for_row(row):
    candidates = []
    for col in ["actual_file_path", "docx_path", "pdf_path", "eml_path", "filename"]:
        if col in row and row[col] and str(row[col]) != "nan":
            candidates.append(str(row[col]))

    if 'filename' in row and row.get('filename') and not any("output" in c for c in candidates):
        candidates.append(os.path.join("output","files", str(row.get("filename"))))

    for p in candidates:
        p = str(p)
        if not p:
            continue
        if not os.path.isabs(p) and not os.path.exists(p):
            p2 = os.path.join("output","files", p)
            if os.path.exists(p2):
                p = p2
        if not os.path.exists(p):
            continue
        lower = p.lower()
        if lower.endswith(".txt"):
            t = read_txt(p)
            if t:
                return t
        elif lower.endswith(".docx"):
            t = read_docx(p)
            if t:
                return t
        elif lower.endswith(".pdf"):
            t = read_pdf(p)
            if t:
                return t
        elif lower.endswith(".eml"):
            t = read_eml(p)
            if t:
                return t
        else:
            t = read_txt(p)
            if t:
                return t
    return ""

def is_placeholder(value):
    if value is None:
        return True
    s = str(value).strip()
    if not s:
        return True
    low = s.lower()

    placeholders = [
        "xxx", "xxxx", "placeholder", "redacted", "example", "sample", "please", "share",
        "confidential", "document", "subject", "generated", "sig=xxxxx", "fake", "n/a",
        "number", "account"
    ]
    for ph in placeholders:
        if ph in low:
            return True

    if "sig=" in low and ("xxxxx" in low or "fake" in low):
        return True

    if re.match(r'^[xX\*\-_]{3,}$', s):
        return True

    if len(set(s)) == 1 and len(s) >= 6:
        return True

    digits_only = re.sub(r"\D", "", s)
    if digits_only:
        if set(digits_only) == {"0"}:
            return True
        if len(digits_only) < 4 and len(digits_only) < len(s):
            return True

    token = re.sub(r"\s+", "", s)
    if len(token) <= 2:
        return True

    if re.match(r"^[a-z0-9._-]{1,6}$", s, flags=re.I):
        return True

    non_alnum = sum(1 for ch in s if not ch.isalnum())
    if non_alnum / max(1, len(s)) > 0.6:
        return True

    if re.search(r'[ilIoO0]{6,}', s):  # long runs of ambiguous OCR chars
        return True

    return False

def normalize_match_obj(m):
    try:
        if hasattr(m, "group"):
            return m.group(0)
        if isinstance(m, (list, tuple)):
            parts = [str(x) for x in m if x is not None]
            return " ".join(parts)
        return str(m)
    except Exception:
        try:
            return str(m)
        except Exception:
            return ""

def sample_for_display(matches, max_items=5):
    out = []
    seen = set()
    for m in matches:
        if isinstance(m, (list, tuple)):
            val = str(m[0]).strip()
            excerpt = m[2] if len(m) > 2 else ""
        else:
            val = str(m).strip()
            excerpt = ""
        if not val:
            continue
        if val in seen:
            continue
        seen.add(val)
        display = val
        if excerpt:
            display = f"{val} ... {excerpt}"
        if len(display) > 120:
            display = display[:117] + "..."
        out.append(display)
        if len(out) >= max_items:
            break
    return out

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    config = load_config()
    df = load_mapping()
    sit_regex = build_sit_regex_map(config)
    per_sit_target = config.get("per_sit_count", 100)

    per_sit_docs = defaultdict(int)
    per_sit_tp_docs = defaultdict(int)
    per_sit_fp_flagged = defaultdict(int)
    per_sit_instance_counts = defaultdict(int)
    issues = []

    import statistics
    sit_counts_per_doc = []
    instance_counts_list = []

    for idx, row in df.iterrows():
        sits = str(row.get("sit_ids",""))
        labels = str(row.get("labels",""))
        instances = str(row.get("instances",""))
        sit_list = [s for s in sits.split(";") if s]
        label_list = [s for s in labels.split(";") if s]
        inst_list = []
        if instances:
            for x in instances.split(";"):
                try:
                    inst_list.append(int(x))
                except Exception:
                    inst_list.append(1)
        text = extract_text_for_row(row)
        sit_counts_per_doc.append(len(sit_list))
        instance_counts_list.extend(inst_list)

        for i, sid in enumerate(sit_list):
            sid = sid.strip()
            per_sit_docs[sid] += 1
            lbl = label_list[i] if i < len(label_list) else "TP"
            inst = inst_list[i] if i < len(inst_list) else 1
            per_sit_instance_counts[sid] += inst

            regex = sit_regex.get(sid)
            matches = []
            if regex and text:
                try:
                    raw_matches = []
                    for m in regex.finditer(text):
                        val = normalize_match_obj(m)
                        start = m.start()
                        end = m.end()
                        excerpt = text[max(0, start-30):min(len(text), end+30)].replace("\n"," ")
                        raw_matches.append((val, start, excerpt))
                    matches = raw_matches
                except Exception:
                    try:
                        raw = regex.findall(text)
                        matches = []
                        pos = 0
                        for r in raw:
                            val = normalize_match_obj(r)
                            excerpt = ""
                            matches.append((val, pos, excerpt))
                            pos += 1
                    except Exception:
                        matches = []

            if lbl == "TP":
                need = max(1, inst)
                # count only non-placeholder matches as "found"
                found_real = 0
                for m in matches:
                    val = m[0] if isinstance(m, (list, tuple)) else m
                    if not is_placeholder(val):
                        found_real += 1
                if found_real >= need:
                    per_sit_tp_docs[sid] += 1
                else:
                    sample_paths = ""
                    try:
                        sample_paths = f"path {row.get('actual_file_path') or row.get('filename')}"
                    except Exception:
                        sample_paths = ""
                    sample = sample_for_display(matches, max_items=5)
                    issues.append(f"TP missing matches for doc {row.get('doc_id')}, sit {sid}: found {found_real} expected {need} {sample_paths} sample_matches: {sample}")
            else:
                valid_found = False
                if matches:
                    for m in matches:
                        val = m[0] if isinstance(m, (list, tuple)) else m
                        if not is_placeholder(val):
                            valid_found = True
                            break
                if valid_found:
                    per_sit_fp_flagged[sid] += 1
                    sample = sample_for_display(matches, max_items=5)
                    issues.append(f"FP contains valid-looking match in doc {row.get('doc_id')}, sit {sid}: sample {sample}")

    lines = []
    lines.append("Validation report\n=================\n")
    lines.append(f"Total unique SITs observed in mapping: {len(per_sit_docs)}\n")

    for sid, count in sorted(per_sit_docs.items()):
        tp = per_sit_tp_docs.get(sid,0)
        fp = per_sit_fp_flagged.get(sid,0)
        inst_sum = per_sit_instance_counts.get(sid,0)
        lines.append(f"{sid}: docs={count}, tp_docs={tp}, fp_flagged={fp}, total_instances={inst_sum}")
        if count < per_sit_target:
            lines.append(f"  >>> WARNING: only {count} docs for {sid} (target {per_sit_target})")

    lines.append("\nDistribution summary:\n")
    if sit_counts_per_doc:
        lines.append(f"Average SITs per doc: {statistics.mean(sit_counts_per_doc):.2f}")
    if instance_counts_list:
        lines.append(f"Average instances per SIT (across docs): {statistics.mean(instance_counts_list):.2f}")

    lines.append("\nDetected issues (first 500 lines):\n")
    if issues:
        lines.extend(issues[:500])
    else:
        lines.append("No issues detected based on regex checks and heuristics.\n")

    with open(REPORT_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print("\n".join(lines[:50]))
    print(f"\nFull report saved to {REPORT_PATH}")

if __name__ == "__main__":
    main()
